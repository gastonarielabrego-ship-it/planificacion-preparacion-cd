# -*- coding: utf-8 -*-
"""Inspección profunda h61 maquinista.xlsx: todas las columnas + distribuciones."""
from collections import Counter
from openpyxl import load_workbook

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"

wb = load_workbook(RUTA, read_only=True, data_only=True)
ws = wb["Datos"]

filas = ws.iter_rows(values_only=True)
enc = next(filas)
print("TODAS LAS COLUMNAS (%d):" % len(enc))
for i, h in enumerate(enc):
    print(f"  col {i+1}: {h!r}")

# acumuladores
func = Counter(); act = Counter(); circ = Counter(); tur = Counter()
fechas = []; ops = set(); filas_n = 0
tot_mue = Counter(); no_cero_ult = Counter()
muestras_ult = {}
tot_mayor = 0; tot_cero = 0

for fila in filas:
    filas_n += 1
    f = list(fila) + [None] * (len(enc) - len(fila))
    d = dict(zip(enc, f))
    func[d.get("FUNCION")] += 1
    act[d.get("ACTIVIDAD")] += 1
    circ[str(d.get("CIRCUITO"))] += 1
    tur[d.get("TURNO")] += 1
    fechas.append(d.get("FECHA"))
    ops.add(d.get("OPERARIO"))
    tm = d.get("TIEMPO_MUE")
    if tm:
        try:
            tot_mue[float(tm)] += 1
        except Exception:
            pass
    # columnas después de HORA_23 (índice de 'TOTAL' en adelante)
    tot = d.get("TOTAL")
    try:
        t = float(tot or 0)
        if t > 0:
            tot_mayor += 1
        else:
            tot_cero += 1
    except Exception:
        pass
    # muestrear valores de las últimas columnas cuando no son cero
    for col in ["TOT_ALMACENAMIENTO", "TOT_APROS", "TOT_APROS_TOTAL", "TOT_APROS_PARCIAL", "TOT_HOMOGENEOS"]:
        v = d.get(col)
        if v not in (None, 0):
            muestras_ult.setdefault(col, Counter())[v] += 1

print("\nFILAS DE DATOS:", filas_n)
print("\nFUNCION:", dict(func))
print("TURNO:", dict(tur))
print("ACTIVIDAD:", dict(act))
print("OPERARIOS distintos:", len(ops))
print("FECHAS: min=%s max=%s | días distintos=%d" % (min(fechas), max(fechas), len(set(fechas))))
print("TOTAL>0: %d | TOTAL=0: %d" % (tot_mayor, tot_cero))
print("\nCIRCUITO (top 40 de %d):" % len(circ))
for k, v in circ.most_common(40):
    print(f"  {k!r}: {v}")
print("\nTIEMPO_MUE no cero (top 10):", tot_mue.most_common(10))
print("\nMuestras TOT_* (valores no cero):")
for col, c in muestras_ult.items():
    print(f"  {col}: {dict(list(c.most_common(5)))}")
wb.close()
print("\nFIN")

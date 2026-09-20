# -*- coding: utf-8 -*-
"""¿Dónde están los APROS y HOMOGENEOS en h61 maquinista.xlsx?"""
from openpyxl import load_workbook
from collections import Counter

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"
wb = load_workbook(RUTA, read_only=True, data_only=True)
print("HOJAS:", wb.sheetnames)

ws = wb[wb.sheetnames[0]]
headers = None
actividad_vals = Counter()
muestra = []
for i, fila in enumerate(ws.iter_rows(values_only=True), start=1):
    vals = list(fila[:45])
    if i == 1:
        headers = vals
        print("\nCOLUMNAS:")
        for j, h in enumerate(headers):
            if h is not None and str(h).strip():
                print(f"  [{j}] {h}")
        continue
    if headers is None:
        continue
    if i <= 6:
        muestra.append(vals)
    # ACTIVIDAD y posibles columnas de funcion/categoria
    if actividad_vals is not None and headers:
        for j, h in enumerate(headers):
            if h is None:
                continue
            hh = str(h).strip().upper()
            if hh in ("ACTIVIDAD", "FUNCION", "CATEGORIA", "TAREA", "DESCRIPCION", "DESCRIP"):
                v = vals[j] if j < len(vals) else None
                if v is not None and str(v).strip():
                    actividad_vals[(hh, str(v).strip()[:40])] += 1

print("\nPRIMERAS FILAS:")
for m in muestra:
    print(" ", [x for x in m if x is not None][:20])

print("\nVALORES DISTINTOS (col actividad/funcion/categoria):")
for (col, v), n in sorted(actividad_vals.items()):
    print(f"  {col} = {v!r}: {n:,} filas")
wb.close()
print("\nFIN")

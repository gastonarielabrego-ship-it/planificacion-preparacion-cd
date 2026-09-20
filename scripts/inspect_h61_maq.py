# -*- coding: utf-8 -*-
"""Inspección de la estructura de h61 maquinista.xlsx (H61 de clarkistas)."""
import sys
from openpyxl import load_workbook

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"

wb = load_workbook(RUTA, read_only=True, data_only=True)
print("HOJAS:", wb.sheetnames)

for nombre in wb.sheetnames:
    ws = wb[nombre]
    print("\n" + "=" * 80)
    print(f"HOJA: {nombre} | dims={ws.calculate_dimension()} | max_row={ws.max_row} max_col={ws.max_column}")
    print("=" * 80)
    n = 0
    for fila in ws.iter_rows(values_only=True):
        n += 1
        # recortar fila a los primeros 40 valores para no ensuciar
        vals = list(fila[:40])
        # compactar: mostrar solo hasta la última celda con dato
        ult = 0
        for i, v in enumerate(vals):
            if v is not None and str(v).strip() != "":
                ult = i
        vals = vals[: ult + 1]
        print(f"  F{n}: {vals}")
        if n >= 25:
            print("  ... (cortado a 25 filas)")
            break
    wb.close() if False else None

# Para la primera hoja: tomar algunas filas del medio y del final para ver variedad
print("\n" + "#" * 80)
print("MUESTREO DEL MEDIO/FINAL (primera hoja)")
print("#" * 80)
ws = wb[wb.sheetnames[0]]
total = ws.max_row
objetivos = sorted(set([total // 2, total - 2, total - 1, total]))
vistos = set()
for i, fila in enumerate(ws.iter_rows(values_only=True), start=1):
    if i in objetivos and i not in vistos:
        vistos.add(i)
        vals = list(fila[:40])
        ult = 0
        for j, v in enumerate(vals):
            if v is not None and str(v).strip() != "":
                ult = j
        print(f"  F{i}: {vals[: ult + 1]}")
wb.close()
print("\nFIN")

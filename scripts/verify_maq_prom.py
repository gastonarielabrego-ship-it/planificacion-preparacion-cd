# -*- coding: utf-8 -*-
"""Sanity check: personas por dia (global) vs por turno en h61 maquinista.xlsx."""
from collections import defaultdict
from openpyxl import load_workbook

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"
wb = load_workbook(RUTA, read_only=True, data_only=True)
ws = wb["Datos"]
filas = ws.iter_rows(values_only=True)
enc = list(next(filas))
idx = {h: i for i, h in enumerate(enc)}

por_dia = defaultdict(set)
por_dia_turno = defaultdict(set)
cruzan = defaultdict(set)
for f in filas:
    fecha = f[idx["FECHA"]]
    op = f[idx["OPERARIO"]]
    tur = f[idx["TURNO"]]
    por_dia[fecha].add(op)
    por_dia_turno[(fecha, tur)].add(op)

tot_dias = len(por_dia)
prom_global = sum(len(v) for v in por_dia.values()) / tot_dias
prom_turno = sum(len(v) for v in por_dia_turno.values()) / len(por_dia_turno)

# personas que aparecen en 2+ turnos el mismo dia
for (fecha, tur), ops in por_dia_turno.items():
    for op in ops:
        cruzan[(fecha, op)].add(tur)
con_cruce = sum(1 for v in cruzan.values() if len(v) > 1)

print("dias:", tot_dias)
print("prom personas/dia (global, sin repetir):", round(prom_global, 1))
print("prom por (dia,turno):", round(prom_turno, 1), "en", len(por_dia_turno), "dia-turnos")
print("persona-dias que cruzan 2+ turnos:", con_cruce)
wb.close()

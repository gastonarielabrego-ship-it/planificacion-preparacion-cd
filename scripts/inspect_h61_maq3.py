# -*- coding: utf-8 -*-
"""Verifica TOT_BULTOS vs TOTAL y horas activas por fila en h61 maquinista.xlsx."""
from openpyxl import load_workbook

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"
wb = load_workbook(RUTA, read_only=True, data_only=True)
ws = wb["Datos"]
filas = ws.iter_rows(values_only=True)
enc = list(next(filas))
idx = {h: i for i, h in enumerate(enc)}

sum_total = 0; sum_bul = 0; sum_bul_alm = 0; n_horas = 0
horas_dist = {}
for f in filas:
    total = f[idx["TOTAL"]] or 0
    bul = f[idx["TOT_BULTOS"]] or 0
    sum_total += total
    sum_bul += bul
    sum_bul_alm += f[idx["TOT_BUL_ALMACENAMIENTO"]] or 0
    nh = sum(1 for h in range(24) if (f[idx[f"HORA_{h:02d}"]] or 0) != 0)
    horas_dist[nh] = horas_dist.get(nh, 0) + 1
    n_horas += nh

print("Σ TOTAL =", sum_total)
print("Σ TOT_BULTOS =", sum_bul)
print("Σ TOT_BUL_ALMACENAMIENTO =", sum_bul_alm)
print("Σ horas activas (por fila) =", n_horas, "→ promedio/fila:", round(n_horas / (horas_dist and sum(horas_dist.values())), 2))
print("Distribución horas activas por fila (top 10):", sorted(horas_dist.items(), reverse=True)[:10])
wb.close()

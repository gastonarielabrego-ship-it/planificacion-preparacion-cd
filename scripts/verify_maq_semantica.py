# -*- coding: utf-8 -*-
"""Relaciones entre columnas TOT_* del archivo maquinistas (apros vs homogeneos)."""
from openpyxl import load_workbook
from collections import Counter

RUTA = "/home/z/my-project/upload/h61 maquinista.xlsx"
wb = load_workbook(RUTA, read_only=True, data_only=True)
ws = wb["Datos"]

idx = {}
tot = Counter()
n = 0
rel_apros_total = Counter()   # APROS_TOTAL ?= APROS + APROS_PARCIAL
rel_homogeneos = Counter()    # HOMOGENEOS ?= HOM_STD + HOM_REALMAC_XD + HOM_REALMAC_STD
rel_total = Counter()         # TOTAL ?= suma HORA_00..23
rel_bultos = Counter()        # TOT_BULTOS ?= APROS_TOTAL + HOMOGENEOS + ALMACENAMIENTO + otros
cross_act = Counter()         # actividad -> filas con apros>0 / homog>0 / ambos / ninguno
suma_apros = Counter()        # actividad -> suma apros
suma_hom = Counter()

for i, fila in enumerate(ws.iter_rows(values_only=True), start=1):
    vals = list(fila[:45])
    if i == 1:
        for j, h in enumerate(vals):
            if h is not None:
                idx[str(h).strip().upper()] = j
        continue
    n += 1
    g = lambda k: (vals[idx[k]] if k in idx and idx[k] < len(vals) and vals[idx[k]] is not None else 0)
    num = lambda k: g(k) if isinstance(g(k), (int, float)) else 0
    act = str(g("ACTIVIDAD"))

    apros = num("TOT_APROS"); apros_t = num("TOT_APROS_TOTAL"); apros_p = num("TOT_APROS_PARCIAL")
    hom = num("TOT_HOMOGENEOS"); hom_std = num("TOT_HOM_STD"); hom_rx = num("TOT_HOM_REALMAC_XD"); hom_rs = num("TOT_HOM_REALMAC_STD")
    total = num("TOTAL"); almac = num("TOT_ALMACENAMIENTO"); bultos = num("TOT_BULTOS")
    horas = sum(num(f"HORA_{h:02d}") for h in range(24))
    tiemue = num("TIEMPO_MUE")

    rel_apros_total["igual" if apros_t == apros + apros_p else f"distinto(t={apros_t},s={apros+apros_p})"] += 1
    rel_homogeneos["igual" if hom == hom_std + hom_rx + hom_rs else f"distinto(h={hom},s={hom_std+hom_rx+hom_rs})"] += 1
    rel_total["igual" if total == horas else f"distinto(t={total},s={horas})"] += 1

    suma = apros_t + hom + almac
    if bultos == suma: rel_bultos["bultos=aprosT+hom+almac"] += 1
    elif bultos == 0: rel_bultos["bultos=0"] += 1
    else: rel_bultos[f"otro(b={bultos},suma={suma},dif={bultos-suma})"] += 1

    tiene_apros = apros_t > 0
    tiene_hom = hom > 0
    cross_act[(act, "apros" if tiene_apros and not tiene_hom else "hom" if tiene_hom and not tiene_apros else "ambos" if tiene_apros and tiene_hom else "nada")] += 1
    suma_apros[act] += apros_t
    suma_hom[act] += hom

print(f"filas: {n:,}")
print("\nTOT_APROS_TOTAL ?= TOT_APROS + TOT_APROS_PARCIAL:")
for k, v in rel_apros_total.most_common(8): print(f"  {k}: {v:,}")
print("\nTOT_HOMOGENEOS ?= HOM_STD + REALMAC_XD + REALMAC_STD:")
for k, v in rel_homogeneos.most_common(8): print(f"  {k}: {v:,}")
print("\nTOTAL ?= Σ HORA_00..23:")
for k, v in rel_total.most_common(8): print(f"  {k}: {v:,}")
print("\nTOT_BULTOS vs (apros_total + homogeneos + almacenamiento):")
for k, v in rel_bultos.most_common(8): print(f"  {k}: {v:,}")
print("\nCruce ACTIVIDAD x tarea (filas):")
for (act, t), v in sorted(cross_act.items()): print(f"  act={act} {t}: {v:,}")
print("\nSumas por actividad: apros / homogeneos:")
for act in sorted(suma_apros): print(f"  act={act}: apros={suma_apros[act]:,} hom={suma_hom[act]:,}")
wb.close()
